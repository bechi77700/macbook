# -*- coding: utf-8 -*-
"""Genere la declaration letter (.docx) et les etats financiers nil (.xlsx)."""
import os, subprocess, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from config import COMPANY, SIGNATORY, TAX_PERIODS, ASSERTIONS, LETTER_DATE

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "out")
os.makedirs(OUT, exist_ok=True)

# ---------------------------------------------------------------- helpers
def _p(doc, text="", size=10, bold=False, align=None, space_after=4, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 0.95
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = "Calibri"
    return p

_MONTHS = {"January": "Jan", "February": "Feb", "March": "Mar", "April": "Apr",
           "May": "May", "June": "Jun", "July": "Jul", "August": "Aug",
           "September": "Sep", "October": "Oct", "November": "Nov", "December": "Dec"}

def _short(d):
    """'1 January 2024' -> '1 Jan 2024' (laisse inchange tout autre format)."""
    for full, abbr in _MONTHS.items():
        if full in d:
            return d.replace(full, abbr)
    return d

# ---------------------------------------------------------------- letter
def build_letter():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(9)
    for s in doc.sections:
        s.top_margin = Cm(1.0); s.bottom_margin = Cm(0.9)
        s.left_margin = Cm(1.9); s.right_margin = Cm(1.9)

    # --- en-tete (supprime ce bloc si tu imprimes sur papier a en-tete)
    _p(doc, COMPANY["name"], size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
    _p(doc, "%s  |  %s" % (COMPANY["legal_form"], ", ".join(COMPANY["address_lines"])),
       size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
    _p(doc, "Trade Licence No. %s (%s)  |  Corporate Tax Ref. No. %s"
       % (COMPANY["licence_no"], COMPANY["licence_authority"], COMPANY["ct_reference"]),
       size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)

    _p(doc, "Date: %s" % LETTER_DATE, size=9, space_after=3)
    _p(doc, "To: The Federal Tax Authority - Corporate Tax Department, United Arab Emirates",
       size=9, bold=True, space_after=4)
    _p(doc, "SUBJECT: DECLARATION OF REVENUES AND ASSETS IN SUPPORT OF CORPORATE TAX "
            "DE-REGISTRATION APPLICATION", size=9, bold=True, space_after=3)

    ref = doc.add_table(rows=0, cols=2)
    ref.alignment = WD_TABLE_ALIGNMENT.LEFT
    rows = [
        ("Taxable Person", COMPANY["name"]),
        ("Corporate Tax Reference No.", COMPANY["ct_reference"]),
        ("Trade Licence No.", "%s, issued by %s" % (COMPANY["licence_no"], COMPANY["licence_authority"])),
        ("Date of incorporation", COMPANY["incorporation_date"]),
        ("Date of cessation of business", COMPANY["cessation_date"]),
    ]
    for k, v in rows:
        c = ref.add_row().cells
        c[0].width = Cm(5.2); c[1].width = Cm(11.9)
        for cell, txt, bold in ((c[0], k, True), (c[1], v, False)):
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            r = cell.paragraphs[0].add_run(txt); r.font.size = Pt(8.5); r.font.bold = bold
    _p(doc, "", space_after=2)

    _p(doc, "Dear Sir / Madam,", size=9, space_after=2)
    _p(doc,
       "Further to the Authority's requests for additional information dated 30 April 2025, "
       "13 February 2026 and 9 September 2026 on the above Corporate Tax de-registration application, "
       "and taking up the option to submit a signed and stamped declaration of revenues and assets in "
       "lieu of financial statements, the Taxable Person declares as follows.", size=9, space_after=3)

    _p(doc, "1.  DECLARATION PER TAX PERIOD", size=9, bold=True, space_after=2)
    _p(doc, "For each and every tax period of the Taxable Person, the figures are as follows:",
       size=9, space_after=2)

    hdr = ["Tax period", "Period covered", "Revenue (AED)", "Total assets (AED)",
           "Total liabilities (AED)", "Taxable income (AED)"]
    t = doc.add_table(rows=1, cols=len(hdr)); t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    colw = [Cm(1.9), Cm(3.8), Cm(2.4), Cm(2.4), Cm(2.5), Cm(2.5)]
    for i, h in enumerate(hdr):
        cell = t.rows[0].cells[i]
        cell.width = colw[i]
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        r = cell.paragraphs[0].add_run(h); r.font.bold = True; r.font.size = Pt(8)
    for tp in TAX_PERIODS:
        cells = t.add_row().cells
        vals = [tp["label"], "%s - %s" % (_short(tp["start"]), _short(tp["end"])),
                "NIL (0.00)", "NIL (0.00)", "NIL (0.00)", "NIL (0.00)"]
        for i, v in enumerate(vals):
            cells[i].width = colw[i]
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)
            r = cells[i].paragraphs[0].add_run(v); r.font.size = Pt(8)
            r.font.bold = (i >= 2)
    _p(doc, "", space_after=2)

    _p(doc, "2.  STATEMENT OF FACTS", size=9, bold=True, space_after=2)
    facts = [
        "The Taxable Person was incorporated on %s under trade licence no. %s issued by %s, and is "
        "not registered for Value Added Tax."
        % (COMPANY["incorporation_date"], COMPANY["licence_no"], COMPANY["licence_authority"]),
        "It never commenced trading. At no point during its existence did it carry on any commercial, "
        "industrial or professional activity, whether within or outside the United Arab Emirates.",
        "It generated no revenue and derived no income of any nature in any tax period; revenue was "
        "AED 0.00 throughout. No taxable income arose and no Corporate Tax is payable for any tax period.",
        "It held no assets and owed no liabilities at any time; total assets and total liabilities were "
        "AED 0.00 at the beginning and at the end of each tax period.",
    ]
    extras = []
    if ASSERTIONS.get("no_bank_account_operated"):
        extras.append("no bank account was opened or operated in its name and no funds were received "
                      "or disbursed by it")
    if ASSERTIONS.get("no_employees"):
        extras.append("it employed no staff and no residence visa was issued under its establishment card")
    if ASSERTIONS.get("no_fixed_assets"):
        extras.append("it owned no fixed assets, inventory, intangible assets or investments at any time")
    if ASSERTIONS.get("setup_costs_borne_by_shareholder"):
        extras.append("such incorporation and licence fees as were incurred were borne personally by the "
                      "shareholder and were never recorded as expenses, liabilities or capital "
                      "contributions of the Taxable Person")
    if extras:
        facts.append("In addition: " + "; ".join(extras) + ".")
    for i, f in enumerate(facts):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 0.95
        p.paragraph_format.left_indent = Cm(0.85)
        p.paragraph_format.first_line_indent = Cm(-0.85)
        r = p.add_run("2.%d  " % (i + 1)); r.font.bold = True; r.font.size = Pt(9)
        r2 = p.add_run(f); r2.font.size = Pt(9)
    _p(doc, "", space_after=2)

    _p(doc, "3.  REASON FOR DE-REGISTRATION (in response to the Authority's query of 30 April 2025)",
       size=9, bold=True, space_after=2)
    if COMPANY["licence_status"] == "cancelled":
        lic = ("its trade licence no. %s was cancelled by %s on %s"
               % (COMPANY["licence_no"], COMPANY["licence_authority"], COMPANY["licence_status_date"]))
    elif COMPANY["licence_status"] == "cancellation_in_progress":
        lic = ("its trade licence no. %s expired on %s, has not been renewed since, and a formal "
               "cancellation application is pending with %s"
               % (COMPANY["licence_no"], COMPANY["licence_status_date"], COMPANY["licence_authority"]))
    else:
        lic = ("its trade licence no. %s expired on %s and has not been renewed at any time since"
               % (COMPANY["licence_no"], COMPANY["licence_status_date"]))
    _p(doc,
       "De-registration is applied for on the ground of cessation of business and business activity. "
       "The Taxable Person ceased all business and business activity on %s and %s. It has no intention "
       "of resuming any activity and exists only as a dormant registration record. The application is "
       "made pursuant to Article 52 of Federal Decree-Law No. 47 of 2022 on the Taxation of "
       "Corporations and Businesses." % (COMPANY["cessation_date"], lic),
       size=9, space_after=4)

    _p(doc,
       "4.  DECLARATION AND UNDERTAKING.  I, the undersigned, being the authorised signatory of the Taxable Person, declare that the "
       "information in this letter is true, complete and accurate in all respects and covers every tax "
       "period of the Taxable Person without exception. I undertake to provide any further document or "
       "clarification the Authority may require.",
       size=9, space_after=4)

    _p(doc, "Yours faithfully,", size=9, space_after=10)
    _p(doc, "_______________________________", size=9, space_after=1)
    _p(doc, SIGNATORY["name"], size=9, bold=True, space_after=1)
    _p(doc, "%s, %s" % (SIGNATORY["title"], COMPANY["name"]), size=8.5, space_after=1)
    _p(doc, "%s: %s   |   Email: %s   |   Mobile: %s"
       % (SIGNATORY["id_label"], SIGNATORY["id_value"], SIGNATORY["email"], SIGNATORY["phone"]),
       size=8.5, space_after=3)
    _p(doc, "[ COMPANY STAMP ]", size=8.5, italic=True, space_after=1)

    path = os.path.join(OUT, "01_Declaration_Letter_FTA_%s.docx" % COMPANY["ct_reference"])
    doc.save(path)
    return path

# ---------------------------------------------------------------- financials
THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HFILL = PatternFill("solid", fgColor="1F3864")
SFILL = PatternFill("solid", fgColor="D9E2F3")

def build_financials():
    wb = Workbook(); wb.remove(wb.active)
    for tp in TAX_PERIODS:
        ws = wb.create_sheet(tp["label"])
        ws.sheet_view.showGridLines = False
        widths = [3, 42, 16, 16, 12]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        r = 1
        def title(txt, size=13, bold=True, fill=None, font_color="000000", height=None):
            nonlocal r
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
            c = ws.cell(row=r, column=2, value=txt)
            c.font = Font(size=size, bold=bold, color=font_color)
            c.alignment = Alignment(horizontal="center", vertical="center")
            if fill:
                for col in range(2, 6):
                    ws.cell(row=r, column=col).fill = fill
            if height: ws.row_dimensions[r].height = height
            r += 1
        def line(label, value=None, bold=False, size=10, indent=0, border=False, fill=None):
            nonlocal r
            c = ws.cell(row=r, column=2, value=("    " * indent) + label)
            c.font = Font(size=size, bold=bold)
            if border: c.border = BOX
            if fill: c.fill = fill
            for col in (3, 4):
                cc = ws.cell(row=r, column=col)
                if border: cc.border = BOX
                if fill: cc.fill = fill
            if value is not None:
                v = ws.cell(row=r, column=4, value=value)
                v.font = Font(size=size, bold=bold)
                v.number_format = '#,##0.00'
                v.alignment = Alignment(horizontal="right")
                if border: v.border = BOX
                if fill: v.fill = fill
            r += 1
        def blank(n=1):
            nonlocal r; r += n

        title(COMPANY["name"], size=14, fill=HFILL, font_color="FFFFFF", height=22)
        title("Trade Licence No. %s  |  %s" % (COMPANY["licence_no"], COMPANY["licence_authority"]),
              size=9, bold=False)
        title("Corporate Tax Reference No. %s" % COMPANY["ct_reference"], size=9, bold=False)
        title("FINANCIAL STATEMENTS - TAX PERIOD %s  (%s to %s)"
              % (tp["label"], tp["start"], tp["end"]), size=11)
        title("All amounts in %s. The company was dormant throughout the period." % COMPANY["currency"],
              size=9, bold=False)
        blank(1)

        # --- 1. Statement of financial position
        title("1.  STATEMENT OF FINANCIAL POSITION (BALANCE SHEET)", size=10, fill=SFILL)
        ws.cell(row=r, column=4, value="As at %s" % tp["end"]).font = Font(size=9, italic=True)
        ws.cell(row=r, column=4).alignment = Alignment(horizontal="right"); r += 1
        line("ASSETS", bold=True)
        line("Non-current assets", indent=1)
        line("Property, plant and equipment", 0, indent=2)
        line("Intangible assets", 0, indent=2)
        line("Total non-current assets", 0, bold=True, indent=1)
        line("Current assets", indent=1)
        line("Inventory", 0, indent=2)
        line("Trade and other receivables", 0, indent=2)
        line("Cash and cash equivalents", 0, indent=2)
        line("Total current assets", 0, bold=True, indent=1)
        line("TOTAL ASSETS", 0, bold=True, border=True)
        blank(1)
        line("EQUITY AND LIABILITIES", bold=True)
        line("Equity", indent=1)
        line("Share capital (issued and unpaid)", 0, indent=2)
        line("Accumulated losses", 0, indent=2)
        line("Total equity", 0, bold=True, indent=1)
        line("Liabilities", indent=1)
        line("Trade and other payables", 0, indent=2)
        line("Due to related parties", 0, indent=2)
        line("Total liabilities", 0, bold=True, indent=1)
        line("TOTAL EQUITY AND LIABILITIES", 0, bold=True, border=True)
        blank(1)

        # --- 2. Statement of profit or loss
        title("2.  STATEMENT OF PROFIT OR LOSS (INCOME STATEMENT)", size=10, fill=SFILL)
        ws.cell(row=r, column=4, value="For the period %s to %s" % (tp["start"], tp["end"])).font = Font(size=9, italic=True)
        ws.cell(row=r, column=4).alignment = Alignment(horizontal="right"); r += 1
        line("Revenue", 0)
        line("Cost of sales", 0)
        line("Gross profit", 0, bold=True)
        line("Administrative and general expenses", 0)
        line("Selling and marketing expenses", 0)
        line("Depreciation and amortisation", 0)
        line("Finance costs", 0)
        line("Other income", 0)
        line("NET PROFIT / (LOSS) FOR THE PERIOD", 0, bold=True, border=True)
        line("Taxable income for Corporate Tax purposes", 0, bold=True, border=True)
        blank(1)

        # --- 3. Trial balance
        title("3.  TRIAL BALANCE", size=10, fill=SFILL)
        for col, lab in ((3, "Debit (AED)"), (4, "Credit (AED)")):
            c = ws.cell(row=r, column=col, value=lab)
            c.font = Font(size=9, bold=True); c.alignment = Alignment(horizontal="right"); c.border = BOX
        c = ws.cell(row=r, column=2, value="Account"); c.font = Font(size=9, bold=True); c.border = BOX
        r += 1
        accounts = ["Cash and cash equivalents", "Trade and other receivables", "Inventory",
                    "Property, plant and equipment", "Trade and other payables",
                    "Due to related parties", "Share capital", "Accumulated losses",
                    "Revenue", "Cost of sales", "Administrative and general expenses"]
        for a in accounts:
            ws.cell(row=r, column=2, value=a).font = Font(size=9.5)
            ws.cell(row=r, column=2).border = BOX
            for col in (3, 4):
                v = ws.cell(row=r, column=col, value=0)
                v.number_format = '#,##0.00'; v.font = Font(size=9.5)
                v.alignment = Alignment(horizontal="right"); v.border = BOX
            r += 1
        ws.cell(row=r, column=2, value="TOTAL").font = Font(size=10, bold=True)
        ws.cell(row=r, column=2).border = BOX
        for col in (3, 4):
            v = ws.cell(row=r, column=col, value=0)
            v.number_format = '#,##0.00'; v.font = Font(size=10, bold=True)
            v.alignment = Alignment(horizontal="right"); v.border = BOX
        r += 2

        # --- notes + signature
        title("NOTES", size=10, fill=SFILL)
        notes = [
            "1.  The company did not commence trading and carried on no business or business "
            "activity at any time during the tax period.",
            "2.  The company generated no revenue and held no assets during the tax period. All "
            "balances above are nil.",
            "3.  These statements have been prepared on a cash basis by management for Corporate "
            "Tax purposes. They are unaudited; the company is not required to prepare audited "
            "financial statements, having revenue below the applicable threshold and not being a "
            "Qualifying Free Zone Person.",
            "4.  These statements are submitted in support of the Corporate Tax de-registration "
            "application under reference %s." % COMPANY["ct_reference"],
        ]
        for n in notes:
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
            c = ws.cell(row=r, column=2, value=n)
            c.font = Font(size=9); c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 26
            r += 1
        blank(1)
        ws.cell(row=r, column=2, value="Signed on behalf of %s" % COMPANY["name"]).font = Font(size=9, bold=True); r += 2
        ws.cell(row=r, column=2, value="_______________________________").font = Font(size=10); r += 1
        ws.cell(row=r, column=2, value=SIGNATORY["name"]).font = Font(size=9, bold=True); r += 1
        ws.cell(row=r, column=2, value=SIGNATORY["title"]).font = Font(size=9); r += 1
        ws.cell(row=r, column=2, value="Date: ____________________        [ COMPANY STAMP ]").font = Font(size=9); r += 1

        ws.print_area = "B1:E%d" % r
        ws.page_setup.orientation = "portrait"
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.4, bottom=0.4)

    path = os.path.join(OUT, "02_Nil_Financial_Statements_%s.xlsx" % COMPANY["ct_reference"])
    wb.save(path)
    return path

def to_pdf(path):
    subprocess.run(["soffice", "--headless", "--norestore", "--convert-to", "pdf",
                    "--outdir", OUT, path], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    return os.path.splitext(path)[0] + ".pdf"

if __name__ == "__main__":
    for f in (build_letter(), build_financials()):
        print(f)
        print(to_pdf(f))
